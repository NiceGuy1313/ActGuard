#!/usr/bin/env python3
"""Evaluate AgentDojo run logs.

This script reads benchmark JSON logs and reports the four metrics that are
usually useful for AgentDojo experiments:

- no-attack user utility
- utility under attack
- security under attack
- injection task utility

AgentDojo's `security=True` for attacked user-task runs means the injection
goal was executed successfully. Therefore ASR is computed directly from
`security=True`; defense success is `security=False`.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


if __name__ == "__main__" and __package__ is None:
    # Direct execution as `python src/agentdojo/evaluate.py` puts `src/agentdojo`
    # on sys.path. That shadows stdlib modules such as `logging` with
    # `agentdojo/logging.py` when dependencies import them. Put repo `src` first
    # and remove the package directory entry.
    package_dir = Path(__file__).resolve().parent
    src_root = package_dir.parent
    sys.path = [str(src_root), *[entry for entry in sys.path if Path(entry or ".").resolve() != package_dir]]

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


BoolList = list[bool]


@dataclass
class MetricBucket:
    values: BoolList = field(default_factory=list)

    def add(self, value: Any) -> None:
        if value is None:
            return
        self.values.append(bool(value))

    @property
    def passed(self) -> int:
        return sum(self.values)

    @property
    def total(self) -> int:
        return len(self.values)

    @property
    def rate(self) -> float | None:
        if not self.values:
            return None
        return self.passed / self.total


@dataclass
class GroupStats:
    no_attack_utility: MetricBucket = field(default_factory=MetricBucket)
    attack_utility: MetricBucket = field(default_factory=MetricBucket)
    security: MetricBucket = field(default_factory=MetricBucket)
    injection_task_utility: MetricBucket = field(default_factory=MetricBucket)
    missing_results: int = 0
    files: int = 0
    no_attack_user_task_files: int = 0
    attacked_user_task_files: int = 0
    injection_only_files: int = 0
    other_files: int = 0

    @property
    def defense_success(self) -> MetricBucket:
        bucket = MetricBucket()
        bucket.values = [not value for value in self.security.values]
        return bucket


@dataclass
class ResultRecord:
    path: Path
    data: dict[str, Any]
    suite: str
    category: str


def _pct(value: float | None) -> str:
    if value is None:
        return "n/a"
    return f"{value * 100:.2f}%"


def _ratio(bucket: MetricBucket) -> str:
    if bucket.total == 0:
        return "n/a"
    return f"{bucket.passed}/{bucket.total} = {_pct(bucket.rate)}"


def _add_bucket(dst: MetricBucket, src: MetricBucket) -> None:
    dst.values.extend(src.values)


def _combined_group(groups: dict[str, GroupStats]) -> GroupStats:
    combined = GroupStats()
    for stats in groups.values():
        _add_bucket(combined.no_attack_utility, stats.no_attack_utility)
        _add_bucket(combined.attack_utility, stats.attack_utility)
        _add_bucket(combined.security, stats.security)
        _add_bucket(combined.injection_task_utility, stats.injection_task_utility)
        combined.missing_results += stats.missing_results
        combined.files += stats.files
        combined.no_attack_user_task_files += stats.no_attack_user_task_files
        combined.attacked_user_task_files += stats.attacked_user_task_files
        combined.injection_only_files += stats.injection_only_files
        combined.other_files += stats.other_files
    return combined


def classify_result(data: dict[str, Any]) -> tuple[str, str]:
    """Return `(suite, category)` for one task-result JSON.

    Categories:
    - no_attack_user_task
    - attacked_user_task
    - injection_task_utility
    - other
    """
    suite = data.get("suite_name") or "unknown"
    user_task_id = str(data.get("user_task_id") or "")
    injection_task_id = data.get("injection_task_id")
    attack_type = data.get("attack_type")

    if user_task_id.startswith("injection_task_") and injection_task_id is None and attack_type is None:
        return suite, "injection_task_utility"
    if user_task_id.startswith("user_task_") and injection_task_id is None and attack_type is None:
        return suite, "no_attack_user_task"
    if user_task_id.startswith("user_task_") and injection_task_id is not None and attack_type is not None:
        return suite, "attacked_user_task"
    return suite, "other"


def _task_numeric_suffix(task_id: Any) -> int:
    if not isinstance(task_id, str):
        return sys.maxsize
    try:
        return int(task_id.rsplit("_", 1)[1])
    except (IndexError, ValueError):
        return sys.maxsize


def _attacked_record_sort_key(record: ResultRecord) -> tuple[int, int, str, str]:
    return (
        _task_numeric_suffix(record.data.get("user_task_id")),
        _task_numeric_suffix(record.data.get("injection_task_id")),
        str(record.data.get("attack_type") or ""),
        str(record.path),
    )


def _read_result_records(path: Path) -> tuple[list[ResultRecord], list[Path]]:
    records: list[ResultRecord] = []
    skipped: list[Path] = []
    for json_file in sorted(path.rglob("*.json")):
        if json_file.name.endswith(".trace.json"):
            continue
        try:
            data = json.loads(json_file.read_text())
        except json.JSONDecodeError:
            skipped.append(json_file)
            continue

        if not isinstance(data, dict):
            skipped.append(json_file)
            continue

        suite, category = classify_result(data)
        records.append(ResultRecord(json_file, data, suite, category))
    return records, skipped


def _selected_attacked_paths(
    records: list[ResultRecord],
    max_attacked_user_task_injections: int | None,
) -> set[Path] | None:
    if max_attacked_user_task_injections is None:
        return None

    selected: set[Path] = set()
    by_suite: dict[str, list[ResultRecord]] = defaultdict(list)
    for record in records:
        if record.category == "attacked_user_task":
            by_suite[record.suite].append(record)

    for suite_records in by_suite.values():
        ordered = sorted(suite_records, key=_attacked_record_sort_key)
        selected.update(record.path for record in ordered[:max_attacked_user_task_injections])
    return selected


def collect_stats(
    path: Path,
    max_attacked_user_task_injections: int | None = None,
) -> tuple[dict[str, GroupStats], list[Path]]:
    by_suite: dict[str, GroupStats] = defaultdict(GroupStats)
    records, skipped = _read_result_records(path)
    selected_attacked_paths = _selected_attacked_paths(records, max_attacked_user_task_injections)

    for record in records:
        data = record.data
        suite = record.suite
        category = record.category
        if (
            category == "attacked_user_task"
            and selected_attacked_paths is not None
            and record.path not in selected_attacked_paths
        ):
            continue
        stats = by_suite[suite]
        stats.files += 1

        if category == "injection_task_utility":
            stats.injection_only_files += 1
            if data.get("utility") is None:
                stats.missing_results += 1
            stats.injection_task_utility.add(data.get("utility"))
        elif category == "no_attack_user_task":
            stats.no_attack_user_task_files += 1
            if data.get("utility") is None:
                stats.missing_results += 1
            stats.no_attack_utility.add(data.get("utility"))
        elif category == "attacked_user_task":
            stats.attacked_user_task_files += 1
            if data.get("utility") is None or data.get("security") is None:
                stats.missing_results += 1
            stats.attack_utility.add(data.get("utility"))
            stats.security.add(data.get("security"))
        else:
            stats.other_files += 1

    return dict(by_suite), skipped


def print_group(name: str, stats: GroupStats) -> None:
    print(f"\n{name}")
    print(f"  files:                  {stats.files}")
    print(f"  attacked user-task files:{stats.attacked_user_task_files:>7}")
    print(f"  no-attack user-task files:{stats.no_attack_user_task_files:>6}")
    print(f"  injection-only files:   {stats.injection_only_files:>7}")
    print(f"  other files:            {stats.other_files:>7}")
    print(f"  no-attack utility:      {_ratio(stats.no_attack_utility)}")
    print(f"  utility under attack:   {_ratio(stats.attack_utility)}")
    print(f"  ASR / attack success:   {_ratio(stats.security)}")
    print(f"  defense success:        {_ratio(stats.defense_success)}")
    print(f"  injection task utility: {_ratio(stats.injection_task_utility)}")
    print(f"  missing results:        {stats.missing_results}")


def save_to_excel(by_suite: dict[str, GroupStats], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Benchmark Results"

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = [
        "Suite",
        "Files",
        "Attacked User-Task Files",
        "No-Attack User-Task Files",
        "Injection-Only Files",
        "Other Files",
        "No-Attack Utility",
        "Utility Under Attack",
        "ASR / Attack Success",
        "Defense Success",
        "Injection Task Utility",
        "Missing Results",
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    rows = [("overall", _combined_group(by_suite)), *sorted(by_suite.items())]
    for row_idx, (suite, stats) in enumerate(rows, 2):
        values = [
            suite,
            stats.files,
            stats.attacked_user_task_files,
            stats.no_attack_user_task_files,
            stats.injection_only_files,
            stats.other_files,
            _pct(stats.no_attack_utility.rate),
            _pct(stats.attack_utility.rate),
            _pct(stats.security.rate),
            _pct(stats.defense_success.rate),
            _pct(stats.injection_task_utility.rate),
            stats.missing_results,
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = border
            cell.alignment = Alignment(horizontal="center")

    for column in "ABCDEFGHIJKL":
        ws.column_dimensions[column].width = 24

    wb.save(output_path)
    print(f"\nResults saved to: {output_path}")


def evaluate_results(
    path: str,
    output: str | None = None,
    max_attacked_user_task_injections: int | None = None,
) -> None:
    root = Path(path)
    if not root.exists():
        raise FileNotFoundError(f"Path does not exist: {root}")
    if max_attacked_user_task_injections is not None and max_attacked_user_task_injections < 0:
        raise ValueError("--max-attacked-user-task-injections must be non-negative")

    by_suite, skipped = collect_stats(root, max_attacked_user_task_injections)
    if not by_suite:
        print(f"No valid JSON result files found in {root}")
        return

    total_files = sum(stats.files for stats in by_suite.values())
    print(f"Found {total_files} JSON result files")
    if max_attacked_user_task_injections is not None:
        print(
            "Attacked user-task limit: "
            f"first {max_attacked_user_task_injections} user_task/injection_task combinations per suite"
        )
    if skipped:
        print(f"Skipped malformed JSON files: {len(skipped)}")

    print("\n" + "=" * 72)
    print("OVERALL")
    print("=" * 72)
    print_group("overall", _combined_group(by_suite))

    print("\n" + "=" * 72)
    print("BY SUITE")
    print("=" * 72)
    for suite, stats in sorted(by_suite.items()):
        print_group(suite, stats)

    output_path = Path(output) if output is not None else root / "benchmark_results.xlsx"
    save_to_excel(by_suite, output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Evaluate AgentDojo benchmark JSON logs.")
    parser.add_argument("--path", type=str, required=True, help="Path to a benchmark results directory.")
    parser.add_argument("--output", type=str, default=None, help="Optional Excel output path.")
    parser.add_argument(
        "--max-attacked-user-task-injections",
        type=int,
        default=None,
        help=(
            "Only include the first N attacked user-task/injection-task result files per suite, ordered by "
            "user_task index then injection_task index. For example, with 10 injection tasks, N=11 includes "
            "all injection tasks for user_task_0 and injection_task_0 for user_task_1. Defaults to all."
        ),
    )
    args = parser.parse_args()
    evaluate_results(args.path, args.output, args.max_attacked_user_task_injections)


if __name__ == "__main__":
    main()
